# SPDX-FileCopyrightText: Copyright (c) 2026 NVIDIA CORPORATION & AFFILIATES. All rights reserved.
# SPDX-License-Identifier: Apache-2.0

"""Evaluate one quantized GR00T server across a LIBERO task suite.

Run this script with the dedicated LIBERO virtual environment. The model server
must already be running. Videos are disabled and results are checkpointed after
every task so long evaluations consume little disk and can be inspected while
they run.
"""

from __future__ import annotations

import argparse
from dataclasses import asdict, dataclass
import json
from pathlib import Path
from typing import Any
import warnings

from libero.libero import benchmark
import numpy as np
from openpyxl import Workbook, load_workbook

from gr00t.eval.rollout_policy import run_gr00t_sim_policy
from gr00t.policy.server_client import PolicyClient


SUITES = ("libero_spatial", "libero_object", "libero_goal", "libero_10")


warnings.filterwarnings("ignore", message=".*Overriding environment .*already in registry.*")
warnings.filterwarnings("ignore", message=".*obs returned by the .* method.*")
warnings.filterwarnings("ignore", message=".*Casting input x to numpy array.*")


@dataclass
class TaskResult:
    task_index: int
    task_name: str
    successes: list[bool]
    success_rate: float
    model1_calls: int
    model2_calls: int
    model2_ratio: float
    episode_lengths: list[int]
    episode_rewards: list[float]


def _parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--suite", choices=SUITES, required=True)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--host", default="127.0.0.1")
    parser.add_argument("--port", type=int, default=5555)
    parser.add_argument("--episodes", type=int, default=10)
    parser.add_argument("--max-episode-steps", type=int, default=720)
    parser.add_argument("--n-action-steps", type=int, default=8)
    parser.add_argument("--seed", type=int, default=42)
    parser.add_argument(
        "--task-indices",
        default="",
        help="Optional comma-separated task indices. Empty evaluates the full suite.",
    )
    parser.add_argument(
        "--baseline",
        type=Path,
        help="Completed baseline JSON for the same suite and episode count.",
    )
    parser.add_argument(
        "--max-extra-failures",
        type=int,
        help="Stop after a task when the candidate has this many more failures than the baseline.",
    )
    return parser.parse_args()


def _json_default(value: Any) -> Any:
    if isinstance(value, np.generic):
        return value.item()
    if isinstance(value, np.ndarray):
        return value.tolist()
    raise TypeError(f"Cannot JSON encode {type(value).__name__}")


def _write_json(path: Path, payload: dict[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary_path = path.with_suffix(path.suffix + ".tmp")
    temporary_path.write_text(
        json.dumps(payload, indent=2, sort_keys=True, default=_json_default) + "\n"
    )
    temporary_path.replace(path)


def _new_routing_workbook(path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Sheet1"
    sheet.append(["model1_calls", "model2_calls", "total_calls", "model1_ratio"])
    workbook.save(path)


def _routing_totals(path: Path, first_row: int) -> tuple[int, int, int]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook["Sheet1"]
    rows = list(sheet.iter_rows(min_row=first_row, values_only=True))
    workbook.close()
    model1_calls = sum(int(row[0] or 0) for row in rows)
    model2_calls = sum(int(row[1] or 0) for row in rows)
    return model1_calls, model2_calls, len(rows)


def main() -> None:
    args = _parse_args()
    if args.episodes <= 0:
        raise ValueError("--episodes must be positive")
    if args.max_extra_failures is not None and args.baseline is None:
        raise ValueError("--max-extra-failures requires --baseline")
    if args.max_extra_failures is not None and args.max_extra_failures < 1:
        raise ValueError("--max-extra-failures must be positive")

    baseline_by_task: dict[int, int] = {}
    if args.baseline is not None:
        with args.baseline.open() as baseline_file:
            baseline_payload = json.load(baseline_file)
        if baseline_payload.get("suite") != args.suite:
            raise ValueError("Baseline suite does not match --suite")
        if baseline_payload.get("episodes_per_task") != args.episodes:
            raise ValueError("Baseline episode count does not match --episodes")
        baseline_by_task = {
            result["task_index"]: sum(bool(value) for value in result["successes"])
            for result in baseline_payload.get("results", baseline_payload.get("tasks", []))
        }

    client = PolicyClient(host=args.host, port=args.port)
    if not client.ping():
        raise RuntimeError(f"Policy server is not reachable at {args.host}:{args.port}")
    del client

    task_names = list(benchmark.get_benchmark_dict()[args.suite]().get_task_names())
    if args.task_indices:
        task_indices = [int(value) for value in args.task_indices.split(",")]
    else:
        task_indices = list(range(len(task_names)))
    invalid_indices = [index for index in task_indices if index < 0 or index >= len(task_names)]
    if invalid_indices:
        raise ValueError(f"Invalid task indices for {args.suite}: {invalid_indices}")

    routing_log_path = args.output.with_suffix(".routing.xlsx")
    _new_routing_workbook(routing_log_path)
    task_results: list[TaskResult] = []
    next_routing_row = 2

    for task_index in task_indices:
        task_name = task_names[task_index]
        print(f"\n[{args.suite}] task {task_index}: {task_name}", flush=True)
        _, successes, episode_infos = run_gr00t_sim_policy(
            env_name=f"libero_sim/{task_name}",
            n_episodes=args.episodes,
            max_episode_steps=args.max_episode_steps,
            policy_client_host=args.host,
            policy_client_port=args.port,
            n_envs=1,
            n_action_steps=args.n_action_steps,
            seed=args.seed + task_index * 1000,
            record_video=False,
            routing_log_path=str(routing_log_path),
            log_inference_latency=False,
        )
        model1_calls, model2_calls, routing_rows = _routing_totals(
            routing_log_path, next_routing_row
        )
        next_routing_row += routing_rows
        total_calls = model1_calls + model2_calls
        result = TaskResult(
            task_index=task_index,
            task_name=task_name,
            successes=[bool(value) for value in successes],
            success_rate=float(np.mean(successes)),
            model1_calls=model1_calls,
            model2_calls=model2_calls,
            model2_ratio=model2_calls / total_calls if total_calls else 0.0,
            episode_lengths=[int(value) for value in episode_infos["episode_lengths"]],
            episode_rewards=[float(value) for value in episode_infos["episode_rewards"]],
        )
        task_results.append(result)

        total_successes = sum(sum(item.successes) for item in task_results)
        total_episodes = sum(len(item.successes) for item in task_results)
        total_model1_calls = sum(item.model1_calls for item in task_results)
        total_model2_calls = sum(item.model2_calls for item in task_results)
        total_route_calls = total_model1_calls + total_model2_calls
        completed_baseline_successes = sum(
            baseline_by_task[result.task_index]
            for result in task_results
            if result.task_index in baseline_by_task
        )
        extra_failures = completed_baseline_successes - total_successes
        payload = {
            "suite": args.suite,
            "seed": args.seed,
            "episodes_per_task": args.episodes,
            "task_indices": task_indices,
            "completed_tasks": len(task_results),
            "successes": total_successes,
            "episodes": total_episodes,
            "success_rate": total_successes / total_episodes,
            "model1_calls": total_model1_calls,
            "model2_calls": total_model2_calls,
            "model2_ratio": total_model2_calls / total_route_calls if total_route_calls else 0.0,
            "tasks": [asdict(item) for item in task_results],
            "baseline_successes_on_completed_tasks": completed_baseline_successes,
            "extra_failures_on_completed_tasks": extra_failures,
            "complete": len(task_results) == len(task_indices),
        }
        _write_json(args.output, payload)
        print(
            f"running success={payload['success_rate']:.3%} "
            f"model2_ratio={payload['model2_ratio']:.3%}",
            flush=True,
        )
        if args.max_extra_failures is not None and extra_failures >= args.max_extra_failures:
            print(
                f"Stopping early: candidate has {extra_failures} extra failures on completed tasks "
                f"(limit {args.max_extra_failures}).",
                flush=True,
            )
            break


if __name__ == "__main__":
    main()
